from pathlib import Path
from datetime import datetime, date
import csv
import math

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from settings import (
    STOCK_RAW_CANDIDATES,
    RAW_STOCK_PARQUET_DIR,
)


# ============================================================
# DataGuide 주가 원본 → 연도별 long Parquet
# 최종 전략에 필요한 필드만 변환합니다.
# ============================================================

ITEM_MAP = {
    "S410000650": ("open_adj", 1.0),
    "S410000700": ("close_adj", 1.0),
    "S410000900": ("trading_value_krw", 1.0),
    # DataGuide 시가총액은 백만원 단위 → 원 단위
    "S410001250": ("market_cap_krw", 1_000_000.0),
}

FEATURES = [
    "open_adj",
    "close_adj",
    "trading_value_krw",
    "market_cap_krw",
]

BUFFER_SIZE = 50_000

RAW_STOCK_PARQUET_DIR.mkdir(parents=True, exist_ok=True)


def resolve_input():
    for p in STOCK_RAW_CANDIDATES:
        if p.exists():
            return p
    msg = "\n".join(str(p) for p in STOCK_RAW_CANDIDATES)
    raise FileNotFoundError(
        "주가 DataGuide 원본을 찾지 못했습니다. settings.py 후보 중 하나를 두세요:\n"
        + msg
    )


def normalize_code(x):
    if x is None:
        return None
    s = str(x).strip()
    if not s:
        return None

    try:
        if isinstance(x, (int, float)):
            if isinstance(x, float) and math.isnan(x):
                return None
            return "A" + str(int(x)).zfill(6)
    except Exception:
        pass

    if s.upper().startswith("A") and s[1:].isdigit():
        return "A" + s[1:].zfill(6)
    if s.isdigit():
        return "A" + s.zfill(6)
    return s


def parse_date(x):
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    if isinstance(x, (int, float)):
        try:
            d = from_excel(x)
            if isinstance(d, datetime):
                return d.date()
            if isinstance(d, date):
                return d
        except Exception:
            pass

    s = str(x).strip()
    for fmt in (
        "%Y%m%d",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_float(x):
    if x is None:
        return None
    if isinstance(x, bool):
        return float(x)
    if isinstance(x, (int, float)):
        try:
            if isinstance(x, float) and math.isnan(x):
                return None
        except Exception:
            pass
        return float(x)

    s = str(x).strip()
    if s in {"", "-", "--", "NA", "N/A", "#N/A", "#VALUE!", "#DIV/0!", "nan", "None"}:
        return None
    s = s.replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except Exception:
        return None


SCHEMA = pa.schema([
    pa.field("date", pa.date32()),
    pa.field("code", pa.string()),
    pa.field("name", pa.string()),
    pa.field("open_adj", pa.float64()),
    pa.field("close_adj", pa.float64()),
    pa.field("trading_value_krw", pa.float64()),
    pa.field("market_cap_krw", pa.float64()),
])

writers = {}
buffers = {}


def get_writer(year):
    if year not in writers:
        out = RAW_STOCK_PARQUET_DIR / f"stock_daily_{year}.parquet"
        if out.exists():
            out.unlink()
        writers[year] = pq.ParquetWriter(out, SCHEMA, compression="zstd")
        buffers[year] = []
    return writers[year]


def flush(year):
    rows = buffers.get(year, [])
    if not rows:
        return
    table = pa.Table.from_pylist(rows, schema=SCHEMA)
    writers[year].write_table(table)
    buffers[year] = []


def text_rows(path):
    last_error = None
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            f = open(path, "r", encoding=enc, newline="")
            header = None
            for line in f:
                cells = line.rstrip("\r\n").split("\t")
                if "코드" in cells and "아이템코드" in cells:
                    header = cells
                    break
            if header is None:
                f.close()
                raise ValueError("DataGuide header를 찾지 못했습니다.")

            reader = csv.reader(f, delimiter="\t")

            def iterator():
                try:
                    for row in reader:
                        yield row
                finally:
                    f.close()

            return header, iterator()
        except UnicodeDecodeError as e:
            last_error = e
            try:
                f.close()
            except Exception:
                pass
            continue
    raise last_error or UnicodeError("파일 인코딩을 판별하지 못했습니다.")


def xlsx_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = None
    header = None
    header_row = None

    for candidate in wb.worksheets:
        for row_no, row in enumerate(
            candidate.iter_rows(min_row=1, max_row=50, values_only=True),
            start=1,
        ):
            vals = [str(x).strip() if x is not None else "" for x in row]
            if "코드" in vals and "아이템코드" in vals:
                ws = candidate
                header = list(row)
                header_row = row_no
                break
        if ws is not None:
            break

    if ws is None:
        wb.close()
        raise ValueError("DataGuide header가 있는 sheet를 찾지 못했습니다.")

    def iterator():
        try:
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                yield list(row)
        finally:
            wb.close()

    return header, iterator()


def write_stock(code, name, feature_arrays, date_cols):
    if code is None or not feature_arrays:
        return 0

    n = 0
    for k, (_, d) in enumerate(date_cols):
        rec = {
            "date": d,
            "code": code,
            "name": name,
            "open_adj": None,
            "close_adj": None,
            "trading_value_krw": None,
            "market_cap_krw": None,
        }
        any_value = False

        for feature in FEATURES:
            vals = feature_arrays.get(feature)
            value = None if vals is None or k >= len(vals) else vals[k]
            rec[feature] = value
            if value is not None:
                any_value = True

        if not any_value:
            continue

        year = d.year
        get_writer(year)
        buffers[year].append(rec)
        n += 1

        if len(buffers[year]) >= BUFFER_SIZE:
            flush(year)

    return n


def main():
    input_file = resolve_input()
    print("Input :", input_file)
    print("Output:", RAW_STOCK_PARQUET_DIR)

    if input_file.suffix.lower() in {".xlsx", ".xlsm"}:
        raw_header, rows = xlsx_rows(input_file)
    else:
        raw_header, rows = text_rows(input_file)

    header = [str(x).strip() if x is not None else "" for x in raw_header]

    idx_code = header.index("코드")
    idx_name = header.index("코드명")
    idx_item = header.index("아이템코드")

    date_cols = []
    for j, v in enumerate(raw_header):
        d = parse_date(v)
        if d is not None:
            date_cols.append((j, d))

    if not date_cols:
        raise ValueError("날짜 컬럼을 찾지 못했습니다.")

    print(f"Dates: {len(date_cols):,} | {date_cols[0][1]} ~ {date_cols[-1][1]}")

    current_code = None
    current_name = None
    arrays = {}
    seen_codes = set()
    found_items = set()
    n_stocks = 0
    n_rows = 0

    for row_no, row in enumerate(rows, start=1):
        raw_code = row[idx_code] if idx_code < len(row) else None
        row_code = normalize_code(raw_code) or current_code
        if row_code is None:
            continue

        if current_code is not None and row_code != current_code:
            n_rows += write_stock(current_code, current_name, arrays, date_cols)
            seen_codes.add(current_code)
            n_stocks += 1
            if row_code in seen_codes:
                raise RuntimeError(
                    f"같은 종목이 분리되어 다시 등장했습니다: {row_code} / row {row_no}"
                )
            arrays = {}
            if n_stocks % 100 == 0:
                print(f"{n_stocks:,} stocks | {n_rows:,} long rows")

        current_code = row_code

        raw_name = row[idx_name] if idx_name < len(row) else None
        if raw_name is not None and str(raw_name).strip():
            current_name = str(raw_name).strip()

        raw_item = row[idx_item] if idx_item < len(row) else None
        if raw_item is None:
            continue
        item_code = str(raw_item).strip()
        if item_code not in ITEM_MAP:
            continue

        found_items.add(item_code)
        feature, scale = ITEM_MAP[item_code]
        vals = []
        for col_idx, _ in date_cols:
            v = to_float(row[col_idx]) if col_idx < len(row) else None
            vals.append(None if v is None else v * scale)
        arrays[feature] = vals

    if current_code is not None:
        n_rows += write_stock(current_code, current_name, arrays, date_cols)
        n_stocks += 1

    for year in list(writers):
        flush(year)
        writers[year].close()

    missing_items = sorted(set(ITEM_MAP) - found_items)

    print("\n" + "=" * 80)
    print("STOCK CONVERSION COMPLETE")
    print("=" * 80)
    print(f"Stocks    : {n_stocks:,}")
    print(f"Long rows : {n_rows:,}")
    print("Found item codes:", sorted(found_items))
    if missing_items:
        print("WARNING missing required item codes:", missing_items)
    print("Output:", RAW_STOCK_PARQUET_DIR)


if __name__ == "__main__":
    main()
