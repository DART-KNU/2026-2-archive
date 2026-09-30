from datetime import datetime, date
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
import pyarrow as pa
import pyarrow.parquet as pq

from settings import CONSENSUS_XLSX, CONSENSUS_DIR

# ============================================================
# DataGuide Consensus wide Excel -> yearly long parquet
# Final strategy에 필요한 두 필드만 저장:
#   FM30041100 -> EPS(Fwd.12M)
#   FM30041515 -> 영업이익(Fwd.12M)
# ============================================================

ITEM_MAP = {
    "FM30041100": "eps_fwd12m",
    "FM30041515": "op_fwd12m",
}

FEATURES = list(ITEM_MAP.values())
BUFFER_SIZE = 50_000

SCHEMA = pa.schema([
    pa.field("date", pa.date32()),
    pa.field("code", pa.string()),
    pa.field("name", pa.string()),
    pa.field("eps_fwd12m", pa.float64()),
    pa.field("op_fwd12m", pa.float64()),
])


def normalize_code(x):
    if x is None:
        return None
    if isinstance(x, float) and math.isnan(x):
        return None

    s = str(x).strip()
    if not s:
        return None

    if isinstance(x, (int, float)):
        return "A" + str(int(x)).zfill(6)
    if s.upper().startswith("A") and s[1:].isdigit():
        return "A" + s[1:].zfill(6)
    if s.isdigit():
        return "A" + s.zfill(6)
    return s


def parse_date(x):
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    if isinstance(x, (int, float)):
        try:
            d = from_excel(x)
            return d.date() if isinstance(d, datetime) else d
        except Exception:
            pass

    s = str(x).strip()
    for fmt in (
        "%Y%m%d",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_float(x):
    if x is None:
        return None
    if isinstance(x, bool):
        return float(x)
    if isinstance(x, (int, float)):
        if isinstance(x, float) and math.isnan(x):
            return None
        return float(x)

    s = str(x).strip()
    if s in {"", "-", "--", "NA", "N/A", "#N/A", "#VALUE!", "#DIV/0!", "nan", "None"}:
        return None
    s = s.replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def find_sheet_and_header(wb):
    for ws in wb.worksheets:
        for row_no, row in enumerate(
            ws.iter_rows(min_row=1, max_row=50, values_only=True),
            start=1,
        ):
            labels = [str(x).strip() if x is not None else "" for x in row]
            if "코드" in labels and "아이템코드" in labels:
                return ws, row_no, list(row), labels
    raise ValueError("DataGuide header가 있는 sheet를 찾지 못했습니다.")


def main():
    if not CONSENSUS_XLSX.exists():
        raise FileNotFoundError(CONSENSUS_XLSX)

    CONSENSUS_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(CONSENSUS_XLSX, read_only=True, data_only=True)
    ws, header_row, raw_header, labels = find_sheet_and_header(wb)

    idx_code = labels.index("코드")
    idx_name = labels.index("코드명") if "코드명" in labels else None
    idx_item = labels.index("아이템코드")

    date_cols = []
    for j, value in enumerate(raw_header):
        d = parse_date(value)
        if d is not None:
            date_cols.append((j, d))

    if not date_cols:
        wb.close()
        raise ValueError("컨센서스 날짜 컬럼을 찾지 못했습니다.")

    writers = {}
    buffers = {}

    def get_writer(year):
        if year not in writers:
            out = CONSENSUS_DIR / f"consensus_{year}.parquet"
            if out.exists():
                out.unlink()
            writers[year] = pq.ParquetWriter(out, SCHEMA, compression="zstd")
            buffers[year] = []
        return writers[year]

    def flush(year):
        rows = buffers.get(year, [])
        if not rows:
            return
        writers[year].write_table(pa.Table.from_pylist(rows, schema=SCHEMA))
        buffers[year] = []

    def write_stock(code, name, arrays):
        n = 0
        for k, (_, d) in enumerate(date_cols):
            rec = {
                "date": d,
                "code": code,
                "name": name,
                "eps_fwd12m": None,
                "op_fwd12m": None,
            }
            any_value = False
            for feature in FEATURES:
                values = arrays.get(feature)
                v = None if values is None or k >= len(values) else values[k]
                rec[feature] = v
                any_value |= v is not None

            if not any_value:
                continue

            year = d.year
            get_writer(year)
            buffers[year].append(rec)
            n += 1
            if len(buffers[year]) >= BUFFER_SIZE:
                flush(year)
        return n

    print("Input :", CONSENSUS_XLSX)
    print("Sheet :", ws.title)
    print(f"Dates : {len(date_cols):,} | {date_cols[0][1]} ~ {date_cols[-1][1]}")

    current_code = None
    current_name = None
    arrays = {}
    n_stocks = 0
    n_long = 0
    found_items = set()

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        row = list(row)
        raw_code = row[idx_code] if idx_code < len(row) else None
        code = normalize_code(raw_code) or current_code
        if code is None:
            continue

        if current_code is not None and code != current_code:
            n_long += write_stock(current_code, current_name, arrays)
            n_stocks += 1
            arrays = {}
            if n_stocks % 100 == 0:
                print(f"{n_stocks:,} stocks | {n_long:,} long rows")

        current_code = code
        if idx_name is not None and idx_name < len(row):
            raw_name = row[idx_name]
            if raw_name is not None and str(raw_name).strip():
                current_name = str(raw_name).strip()

        raw_item = row[idx_item] if idx_item < len(row) else None
        if raw_item is None:
            continue
        item_code = str(raw_item).strip()
        if item_code not in ITEM_MAP:
            continue

        found_items.add(item_code)
        feature = ITEM_MAP[item_code]
        arrays[feature] = [
            to_float(row[j]) if j < len(row) else None
            for j, _ in date_cols
        ]

    if current_code is not None:
        n_long += write_stock(current_code, current_name, arrays)
        n_stocks += 1

    wb.close()

    for year in list(writers):
        flush(year)
        writers[year].close()

    missing = sorted(set(ITEM_MAP) - found_items)
    if missing:
        raise ValueError(f"필수 컨센서스 item code 누락: {missing}")

    print("\nCONSENSUS CONVERSION COMPLETE")
    print(f"Stocks   : {n_stocks:,}")
    print(f"Long rows: {n_long:,}")
    print("Output   :", CONSENSUS_DIR)


if __name__ == "__main__":
    main()
